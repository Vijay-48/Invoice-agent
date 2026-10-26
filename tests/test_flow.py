"""API + worker + Excel, end to end with a fake extractor."""
import threading
from datetime import datetime, timedelta, timezone
from pathlib import Path
from types import SimpleNamespace

import pymupdf
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import update

from app import worker, excel
from app.api import app
from app.config import settings
from app.db import SessionLocal
from app.ingest import ingest_document
from app.models import Invoice, Status
from app.pipeline import extract
from tests.conftest import make_extracted, not_an_invoice, SUPPLIER_GSTIN


@pytest.fixture
def client():
    with TestClient(app) as c:
        yield c


@pytest.fixture
def fake_extract(monkeypatch):
    """Each call pops the next ExtractedInvoice (or exception) from the queue."""
    queue = []

    def _extract(path, name=None, **kw):
        item = queue.pop(0)
        if isinstance(item, Exception):
            raise item
        return item

    monkeypatch.setattr(worker, "extract_invoice", _extract)
    monkeypatch.setattr(worker, "page_count", lambda path: 1)   # uploads here aren't real PDFs
    return queue


def h(user):
    return {"X-User-Id": str(user.id)}


def upload(client, user, content=b"%PDF-1.4 fake", name="inv.pdf"):
    r = client.post("/api/v1/invoices/upload", files={"file": (name, content, "application/pdf")}, headers=h(user))
    assert r.status_code == 201, r.text
    return r.json()


def run_worker():
    while (inv_id := worker.claim_next()) is not None:
        worker.process(inv_id)
    worker.export()


def get(client, inv_id):
    return client.get(f"/api/v1/invoices/{inv_id}").json()


FY = "2026-27"   # make_extracted's invoices are dated 2026-09-15


def excel_rows(sheet="Invoices", fy=FY):
    wb = load_workbook(excel.workbook_path(fy))
    return [list(r) for r in wb[sheet].iter_rows(min_row=2, values_only=True)]


# ---------- users

def test_users(client):
    r = client.post("/api/v1/users", json={"name": "Priya", "email": "Priya@Example.com"})
    assert r.status_code == 201 and r.json()["email"] == "priya@example.com"
    assert r.json()["gmail"]["connected"] is False
    assert client.post("/api/v1/users", json={"name": "P2", "email": "priya@example.com"}).status_code == 409
    assert client.post("/api/v1/users", json={"name": "X", "email": "not-an-email"}).status_code == 422
    assert [u["name"] for u in client.get("/api/v1/users").json()] == ["Priya"]
    assert client.get("/api/v1/users/99").status_code == 404


# ---------- upload

def test_upload_validation(client, users):
    priya = users["priya"]
    assert client.post("/api/v1/invoices/upload", files={"file": ("a.gif", b"x")}, headers=h(priya)).status_code == 400
    assert client.post("/api/v1/invoices/upload", files={"file": ("a.pdf", b"x")}).status_code == 422
    assert client.post("/api/v1/invoices/upload", files={"file": ("a.pdf", b"x")},
                       headers={"X-User-Id": "99"}).status_code == 401
    first = upload(client, priya, b"same bytes")
    assert (first["invoice_id"], first["status"], first["duplicate_file"]) == (1, Status.RECEIVED, False)
    again = upload(client, users["ravi"], b"same bytes")
    assert (again["invoice_id"], again["duplicate_file"]) == (1, True)
    assert client.get("/api/v1/invoices/999").status_code == 404


def test_concurrent_same_file_ingest(users):
    results, errors = [], []
    barrier = threading.Barrier(4)

    def go():
        try:
            with SessionLocal() as s:
                barrier.wait()
                inv, created = ingest_document(s, b"racing bytes", "r.pdf", "upload", users["priya"].id)
                results.append((inv.id, created))
        except Exception as e:  # noqa: BLE001
            errors.append(e)

    threads = [threading.Thread(target=go) for _ in range(4)]
    [t.start() for t in threads]
    [t.join() for t in threads]
    assert not errors, errors
    assert len({i for i, _ in results}) == 1 and sum(c for _, c in results) == 1


# ---------- outcomes

def test_correct_invoice_goes_straight_to_excel(client, users, fake_extract):
    fake_extract.append(make_extracted())
    inv_id = upload(client, users["priya"])["invoice_id"]
    run_worker()

    d = get(client, inv_id)
    assert d["status"] == Status.EXPORTED and d["approved_by"] == "auto", d
    assert [a["stage"] for a in d["audit"]] == ["ingest", "extract", "decide", "export"]

    [row] = excel_rows()
    assert row[:5] == [inv_id, "Sri Lakshmi Steels Pvt Ltd", SUPPLIER_GSTIN, "SLS/26-27/118",
                       datetime(2026, 9, 15)]
    assert row[13] == 6136 and row[14] == "auto" and row[16] == "upload" and row[18] == "Priya"
    lines = excel_rows("Line Items")
    assert [(l[0], l[3], l[4], l[10]) for l in lines] == [(inv_id, 1, "MS Pipe 25mm", 810),
                                                          (inv_id, 2, "GI Elbow 25mm", 126)]
    r = client.get("/api/v1/excel")
    assert r.status_code == 200 and r.content[:2] == b"PK"


def test_needs_review_then_correct_and_approve(client, users, fake_extract):
    fake_extract.append(make_extracted(grand_total=6163, confidence=0.6, low_confidence_fields=["grand_total"]))
    inv_id = upload(client, users["priya"])["invoice_id"]
    run_worker()
    d = get(client, inv_id)
    assert d["status"] == Status.NEEDS_REVIEW and "Low extraction confidence" in d["status_reason"]
    assert not excel.list_workbooks()

    ravi = users["ravi"]
    assert client.patch(f"/api/v1/invoices/{inv_id}", json={}, headers=h(ravi)).status_code == 400
    r = client.patch(f"/api/v1/invoices/{inv_id}", json={"grand_total": 6136, "note": "typo"}, headers=h(ravi))
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["status"] == Status.NEEDS_REVIEW and "Ready to approve" in d["status_reason"]
    assert all(c["status"] == "PASS" for c in d["checks"])

    r = client.post(f"/api/v1/invoices/{inv_id}/approve", json={"note": "checked"}, headers=h(ravi))
    assert r.status_code == 200 and r.json()["status"] == Status.EXPORTED and r.json()["excel_error"] is None
    [row] = excel_rows()
    assert row[13] == 6136 and row[14] == "Ravi"
    stages = [(a["stage"], a["actor"]) for a in get(client, inv_id)["audit"]]
    assert ("correct", "Ravi") in stages and ("review", "Ravi") in stages
    assert client.post(f"/api/v1/invoices/{inv_id}/approve", headers=h(ravi)).status_code == 409


def test_over_limit_needs_approval(client, users, fake_extract):
    lines = [dict(description="Steel coil", hsn_sac="7208", quantity=1, unit_price=60000,
                  taxable_value=60000, gst_rate=18)]
    fake_extract.append(make_extracted(lines=lines))
    inv_id = upload(client, users["priya"])["invoice_id"]
    run_worker()
    d = get(client, inv_id)
    assert d["status"] == Status.NEEDS_REVIEW and "auto-approve limit" in d["status_reason"]


def test_not_invoice_and_reject(client, users, fake_extract):
    fake_extract.append(not_an_invoice())
    inv_id = upload(client, users["priya"], name="terms.pdf")["invoice_id"]
    run_worker()
    d = get(client, inv_id)
    assert d["status"] == Status.NOT_INVOICE and d["checks"] == []
    assert client.post(f"/api/v1/invoices/{inv_id}/approve", headers=h(users["priya"])).status_code == 409
    r = client.post(f"/api/v1/invoices/{inv_id}/reject", json={"note": "T&C"}, headers=h(users["priya"]))
    assert r.json()["status"] == Status.REJECTED


def test_duplicate_invoice_from_another_channel(client, users, fake_extract):
    fake_extract += [make_extracted(), make_extracted(invoice_number="SLS 26-27 118")]
    first = upload(client, users["priya"], b"pdf from email")["invoice_id"]
    second = upload(client, users["ravi"], b"phone photo of same invoice")["invoice_id"]
    run_worker()
    assert get(client, first)["status"] == Status.EXPORTED
    d = get(client, second)
    assert d["status"] == Status.DUPLICATE and d["duplicate_of"] == first
    assert len(excel_rows()) == 1


def test_reprocess_after_misread(client, users, fake_extract):
    fake_extract += [not_an_invoice(), make_extracted()]
    inv_id = upload(client, users["priya"])["invoice_id"]
    run_worker()
    assert get(client, inv_id)["status"] == Status.NOT_INVOICE
    r = client.post(f"/api/v1/invoices/{inv_id}/reprocess", headers=h(users["priya"]))
    assert r.json()["status"] == Status.RECEIVED
    run_worker()
    assert get(client, inv_id)["status"] == Status.EXPORTED


def test_retries_then_failed(client, users, fake_extract):
    fake_extract += [RuntimeError("openai down")] * worker.MAX_ATTEMPTS
    inv_id = upload(client, users["priya"])["invoice_id"]
    run_worker()
    d = get(client, inv_id)
    assert d["status"] == Status.FAILED and "openai down" in d["error"]


def test_excel_open_in_excel_waits_and_retries(client, users, fake_extract, monkeypatch):
    real_replace = excel.os.replace

    def locked(src, dst):
        raise PermissionError("file is open")

    monkeypatch.setattr(excel.os, "replace", locked)
    fake_extract.append(make_extracted())
    inv_id = upload(client, users["priya"])["invoice_id"]
    run_worker()
    assert get(client, inv_id)["status"] == Status.APPROVED     # waiting, not lost
    assert not list(Path(settings.excel_path).parent.glob("*.tmp.xlsx"))

    monkeypatch.setattr(excel.os, "replace", real_replace)
    worker.export()
    assert get(client, inv_id)["status"] == Status.EXPORTED
    worker.export()                                              # nothing new: no duplicate rows
    assert len(excel_rows()) == 1


def test_export_appends_to_existing_workbook_without_duplicates(client, users, fake_extract):
    fake_extract += [make_extracted(), make_extracted(invoice_number="SLS/26-27/119")]
    upload(client, users["priya"], b"one")
    run_worker()
    second = upload(client, users["priya"], b"two")["invoice_id"]
    # Simulate a crash after saving but before marking EXPORTED: the row is already there
    run_worker()
    with SessionLocal() as s:
        s.execute(update(Invoice).where(Invoice.id == second).values(status=Status.APPROVED))
        s.commit()
    worker.export()
    assert [r[0] for r in excel_rows()] == [1, second]


def test_stale_processing_is_reclaimed(client, users, fake_extract):
    inv_id = upload(client, users["priya"])["invoice_id"]
    assert worker.claim_next() == inv_id          # claimed, then the worker "crashes"
    assert worker.claim_next() is None
    with SessionLocal() as s:
        s.execute(update(Invoice).where(Invoice.id == inv_id).values(
            updated_at=datetime.now(timezone.utc) - timedelta(hours=1)))
        s.commit()
    fake_extract.append(make_extracted())
    run_worker()
    assert get(client, inv_id)["status"] == Status.EXPORTED


def test_too_many_pages_needs_review(client, users, fake_extract, monkeypatch):
    monkeypatch.setattr(worker, "page_count", lambda path: 9)
    fake_extract.append(make_extracted())
    inv_id = upload(client, users["priya"])["invoice_id"]
    run_worker()
    d = get(client, inv_id)
    assert d["status"] == Status.NEEDS_REVIEW and "first 5 of 9 pages" in d["status_reason"]


def test_list_filters_and_file_download(client, users, fake_extract):
    fake_extract += [make_extracted(), not_an_invoice()]
    upload(client, users["priya"], b"a")
    upload(client, users["ravi"], b"b")
    run_worker()
    assert [i["status"] for i in client.get("/api/v1/invoices?status=exported").json()] == [Status.EXPORTED]
    assert len(client.get(f"/api/v1/invoices?user_id={users['ravi'].id}").json()) == 1
    assert client.get("/api/v1/invoices/1/file").content == b"a"


def test_parallel_duplicates_only_one_passes(client, users, monkeypatch):
    barrier = threading.Barrier(2)
    monkeypatch.setattr(worker, "extract_invoice", lambda path, name=None, **kw: (barrier.wait(timeout=10), make_extracted())[1])
    monkeypatch.setattr(worker, "page_count", lambda path: 1)
    ids = [upload(client, users["priya"], f"copy {i}".encode())["invoice_id"] for i in range(2)]
    claimed = [worker.claim_next() for _ in ids]
    threads = [threading.Thread(target=worker.process, args=(i,)) for i in claimed]
    [t.start() for t in threads]
    [t.join() for t in threads]
    assert sorted(get(client, i)["status"] for i in ids) == [Status.APPROVED, Status.DUPLICATE]


# ---------- extract (no OpenAI call)

def test_prepare_pages_and_extract(tmp_path, monkeypatch):
    pdf = tmp_path / "a.pdf"
    doc = pymupdf.open()
    for i in range(7):
        doc.new_page().insert_text((72, 72), f"TAX INVOICE page {i + 1} total 6136.00")
    doc.save(pdf)
    doc.close()
    images, text, total = extract.prepare_pages(str(pdf))
    # Over the limit: the first and last pages, then the next in order when all look alike
    assert len(images) == 5 and "page 7 ---" in text and "page 4 ---" in text and "page 5 ---" not in text
    assert extract.page_count(str(pdf)) == 7

    seen = {}

    def fake_run(prompt, images):
        seen["prompt"], seen["n"] = prompt, len(images)
        return SimpleNamespace(content=make_extracted())

    monkeypatch.setattr(extract.extractor, "run", fake_run)
    assert extract.extract_invoice(str(pdf)).invoice_number == "SLS/26-27/118"
    assert seen["n"] == 5 and "6136.00" in seen["prompt"]

    monkeypatch.setattr(extract.extractor, "run", lambda *a, **k: SimpleNamespace(content="Incorrect API key"))
    with pytest.raises(ValueError):
        extract.extract_invoice(str(pdf))


def test_invoice_the_ai_called_other_goes_to_review(client, users, fake_extract):
    """Real case: a scanned 'ITI Inv no. 106.pdf' was labelled 'other' and silently dropped."""
    fake_extract += [make_extracted(document_type="other"), not_an_invoice()]
    scan = upload(client, users["priya"], b"scanned invoice")["invoice_id"]
    terms = upload(client, users["priya"], b"terms and conditions")["invoice_id"]
    run_worker()
    d = get(client, scan)
    assert d["status"] == Status.NEEDS_REVIEW and "looks like an invoice" in d["status_reason"]
    assert d["checks"]                                  # validated, so the reviewer sees the checks
    assert get(client, terms)["status"] == Status.NOT_INVOICE


def test_scanned_pdf_uses_scan_model(tmp_path, monkeypatch):
    scan = tmp_path / "scan.pdf"
    doc = pymupdf.open()
    doc.new_page()                                      # image-only page: no text layer
    doc.save(scan)
    doc.close()
    used = []
    for name, agent in [("text", extract.extractor), ("scan", extract.scan_extractor)]:
        monkeypatch.setattr(agent, "run", lambda prompt, images, n=name: (
            used.append((n, "Text layer" in prompt)), SimpleNamespace(content=make_extracted()))[1])
    extract.extract_invoice(str(scan), "ITI Inv no. 106.pdf")
    assert used[-1] == ("scan", False)


# ---------- journal posting entries

def journal(**kw):
    """Tally journal voucher No. 075 recording invoice SLS/26-27/118 (no GSTINs printed)."""
    return make_extracted(**{"document_type": "journal_entry", "supplier_gstin": None, "buyer_gstin": None,
                             "voucher_number": "075", "voucher_date": "2026-09-20", **kw})


def test_journal_after_its_invoice_is_duplicate(client, users, fake_extract):
    fake_extract += [make_extracted(), journal()]
    inv = upload(client, users["priya"], b"invoice")["invoice_id"]
    run_worker()
    j = upload(client, users["priya"], b"journal")["invoice_id"]
    run_worker()
    assert get(client, inv)["status"] == Status.EXPORTED
    d = get(client, j)
    assert d["status"] == Status.DUPLICATE and d["duplicate_of"] == inv
    assert "Journal entry for invoice" in d["status_reason"]
    assert len(excel_rows()) == 1                       # the purchase is counted once


def test_journal_before_its_invoice_is_linked_when_invoice_arrives(client, users, fake_extract):
    fake_extract += [journal(), make_extracted()]
    j = upload(client, users["priya"], b"journal")["invoice_id"]
    run_worker()
    d = get(client, j)
    assert d["status"] == Status.NEEDS_REVIEW and "hasn't been received yet" in d["status_reason"]
    inv = upload(client, users["priya"], b"invoice")["invoice_id"]
    run_worker()
    assert get(client, inv)["status"] == Status.EXPORTED
    assert get(client, j)["status"] == Status.DUPLICATE and get(client, j)["duplicate_of"] == inv


def test_invoice_after_approved_journal_needs_review(client, users, fake_extract):
    fake_extract += [journal(), make_extracted()]
    j = upload(client, users["priya"], b"journal")["invoice_id"]
    run_worker()
    client.post(f"/api/v1/invoices/{j}/approve", headers=h(users["priya"]))
    assert excel_rows()[0][20] == "Journal entry"
    inv = upload(client, users["priya"], b"invoice")["invoice_id"]
    run_worker()
    d = get(client, inv)
    assert d["status"] == Status.NEEDS_REVIEW and "count the purchase twice" in d["status_reason"]


def test_process_as_invoice_override(client, users, fake_extract):
    fake_extract.append(make_extracted(document_type="other", supplier_gstin=None))   # AI unsure, no GSTIN
    inv = upload(client, users["priya"])["invoice_id"]
    run_worker()
    assert get(client, inv)["status"] == Status.NOT_INVOICE
    r = client.post(f"/api/v1/invoices/{inv}/process-as-invoice", json={"note": "it is an invoice"},
                    headers=h(users["ravi"]))
    assert r.status_code == 200
    d = r.json()
    assert d["status"] == Status.NEEDS_REVIEW and any(c["name"] == "supplier_gstin" for c in d["checks"])
    client.patch(f"/api/v1/invoices/{inv}", json={"supplier_gstin": SUPPLIER_GSTIN}, headers=h(users["ravi"]))
    assert client.post(f"/api/v1/invoices/{inv}/approve", headers=h(users["ravi"])).json()["status"] == Status.EXPORTED
    assert excel_rows()[0][20] == "Invoice"
    assert client.post(f"/api/v1/invoices/{inv}/process-as-invoice", headers=h(users["ravi"])).status_code == 409


def test_old_workbook_gets_document_column(client, users, fake_extract):
    from openpyxl import Workbook
    Path(settings.excel_path).parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.active.title = "Invoices"
    wb["Invoices"].append([c[0] for c in excel.INVOICE_COLUMNS[:-1]])   # header from the previous version
    wb["Invoices"].append([99] + [None] * 19)
    wb.save(excel.workbook_path(FY))
    fake_extract.append(make_extracted())
    upload(client, users["priya"])
    run_worker()
    ws = load_workbook(excel.workbook_path(FY))["Invoices"]
    assert ws.cell(1, 21).value == "Document" and ws.cell(3, 21).value == "Invoice" and ws.cell(2, 1).value == 99


# ---------- review: fixes from the full-flow audit

def test_reject_reopens_its_duplicates(client, users, fake_extract):
    fake_extract += [make_extracted(grand_total=1, confidence=0.5), make_extracted(grand_total=1, confidence=0.5),
                     journal(grand_total=1)]
    first = upload(client, users["priya"], b"copy one")["invoice_id"]
    run_worker()
    copy = upload(client, users["priya"], b"copy two")["invoice_id"]
    jrnl = upload(client, users["priya"], b"journal")["invoice_id"]
    run_worker()
    assert get(client, copy)["status"] == Status.DUPLICATE and get(client, jrnl)["status"] == Status.DUPLICATE
    client.post(f"/api/v1/invoices/{first}/reject", json={"note": "misread"}, headers=h(users["ravi"]))
    for i in (copy, jrnl):
        d = get(client, i)
        assert d["status"] == Status.NEEDS_REVIEW and d["duplicate_of"] is None
        assert f"duplicate of #{first}, which Ravi rejected" in d["status_reason"]


def test_second_copy_of_a_journal_is_duplicate(client, users, fake_extract):
    fake_extract += [journal(), journal()]
    first = upload(client, users["priya"], b"journal")["invoice_id"]
    second = upload(client, users["priya"], b"journal sent again")["invoice_id"]
    run_worker()
    assert get(client, first)["status"] == Status.NEEDS_REVIEW
    d = get(client, second)
    assert d["status"] == Status.DUPLICATE and d["duplicate_of"] == first


def test_rejected_can_be_reprocessed(client, users, fake_extract):
    fake_extract += [not_an_invoice(), make_extracted()]
    inv = upload(client, users["priya"])["invoice_id"]
    run_worker()
    client.post(f"/api/v1/invoices/{inv}/reject", headers=h(users["priya"]))
    assert client.post(f"/api/v1/invoices/{inv}/reprocess", headers=h(users["priya"])).json()["status"] == Status.RECEIVED
    run_worker()
    assert get(client, inv)["status"] == Status.EXPORTED


def test_input_validation(client, users):
    assert client.post("/api/v1/invoices/upload", files={"file": ("a.pdf", b"")},
                       headers=h(users["priya"])).status_code == 400
    assert client.get("/api/v1/invoices?limit=0").status_code == 422
    assert client.get("/api/v1/invoices?offset=-1").status_code == 422
    assert client.post("/api/v1/users", json={"name": "  ", "email": "x@y.com"}).status_code == 422
    assert client.post(f"/api/v1/users/{users['priya'].id}/gmail/sync?days=0").status_code == 422


def test_excel_open_is_shown_on_waiting_invoices_and_logged_once(client, users, fake_extract, monkeypatch, caplog):
    monkeypatch.setattr(excel.os, "replace", lambda s, d: (_ for _ in ()).throw(PermissionError("open")))
    monkeypatch.setattr(excel, "_last_problem", None)
    fake_extract.append(make_extracted())
    inv = upload(client, users["priya"])["invoice_id"]
    run_worker()
    worker.export()
    worker.export()
    d = get(client, inv)
    assert d["status"] == Status.APPROVED and "open in another program" in d["error"]
    assert sum("open in another program" in r.message for r in caplog.records if r.name == "excel") == 1


def test_audit_records_model_used(client, users, fake_extract):
    data = make_extracted()
    data._model = "gpt-4o"
    fake_extract.append(data)
    inv = upload(client, users["priya"])["invoice_id"]
    run_worker()
    extract_audit = next(a for a in get(client, inv)["audit"] if a["stage"] == "extract")
    assert extract_audit["data"]["model"] == "gpt-4o"



def test_journal_amounts_are_checked(client, users, fake_extract):
    fake_extract.append(journal(grand_total=6500))          # party credit doesn't match the debits
    j = upload(client, users["priya"], b"journal")["invoice_id"]
    run_worker()
    d = get(client, j)
    assert d["status"] == Status.NEEDS_REVIEW and "expected 6136.00, got 6500.00" in d["status_reason"]


def test_approved_journal_in_excel_with_voucher(client, users, fake_extract):
    fake_extract.append(journal())
    j = upload(client, users["priya"], b"journal")["invoice_id"]
    run_worker()
    client.post(f"/api/v1/invoices/{j}/approve", headers=h(users["priya"]))
    row = excel_rows()[0]
    assert row[3] == "SLS/26-27/118" and row[20:23] == ["Journal entry", "075", datetime(2026, 9, 20)]
    assert len(excel_rows("Line Items")) == 2


def test_correcting_data_stored_before_voucher_fields_existed(client, users, fake_extract):
    fake_extract.append(make_extracted(grand_total=1))
    inv = upload(client, users["priya"])["invoice_id"]
    run_worker()
    with SessionLocal() as s:                                 # as stored by the previous version
        row = s.get(Invoice, inv)
        row.extracted = {k: v for k, v in row.extracted.items()
                         if k not in ("voucher_number", "voucher_date", "invoice_kind",
                                      "reference_invoice_number", "currency")}
        s.commit()
    r = client.patch(f"/api/v1/invoices/{inv}", json={"grand_total": 6136}, headers=h(users["priya"]))
    assert r.status_code == 200 and "Ready to approve" in r.json()["status_reason"]



# ---------- invoice types

def test_credit_note_is_reviewed_and_negative_in_excel(client, users, fake_extract):
    fake_extract.append(make_extracted(document_type="credit_note", invoice_number="CN-7",
                                       reference_invoice_number="SLS/26-27/118"))
    cn = upload(client, users["priya"], b"credit note")["invoice_id"]
    run_worker()
    d = get(client, cn)
    assert d["status"] == Status.NEEDS_REVIEW
    assert d["status_reason"].startswith("Credit note against invoice SLS/26-27/118: reduces what you owe")
    client.post(f"/api/v1/invoices/{cn}/approve", headers=h(users["priya"]))
    row = excel_rows()[0]
    assert row[13] == -6136 and row[20] == "Credit note" and row[24] == "SLS/26-27/118"
    lines = excel_rows("Line Items")
    assert (lines[0][6], lines[0][8], lines[0][10]) == (-10, -4500, -810)   # qty x rate still = taxable


def test_kinds_that_can_repeat_amounts_go_to_review(client, users, fake_extract):
    cases = {"payment_reminder": "Payment reminder", "consolidated": "Consolidated invoice",
             "final": "Final project invoice"}
    for i, kind in enumerate(cases):
        fake_extract.append(make_extracted(invoice_kind=kind, invoice_number=f"K-{i}"))
        upload(client, users["priya"], kind.encode())
    run_worker()
    for i, (kind, text) in enumerate(cases.items(), start=1):
        d = get(client, i)
        assert d["status"] == Status.NEEDS_REVIEW and d["status_reason"].startswith(text), (kind, d)


def test_reminder_for_recorded_invoice_is_duplicate(client, users, fake_extract):
    fake_extract += [make_extracted(), make_extracted(invoice_kind="payment_reminder")]
    first = upload(client, users["priya"], b"invoice")["invoice_id"]
    run_worker()
    reminder = upload(client, users["priya"], b"overdue reminder")["invoice_id"]
    run_worker()
    assert get(client, reminder)["status"] == Status.DUPLICATE and get(client, reminder)["duplicate_of"] == first


def test_foreign_currency_needs_review(client, users, fake_extract):
    fake_extract.append(make_extracted(currency="usd", invoice_kind="import"))
    inv = upload(client, users["priya"])["invoice_id"]
    run_worker()
    assert "Amounts are in USD, not INR" in get(client, inv)["status_reason"]


def test_regular_kinds_auto_approve_and_show_type(client, users, fake_extract):
    fake_extract.append(make_extracted(invoice_kind="recurring", irn="a" * 64))
    upload(client, users["priya"])
    run_worker()
    row = excel_rows()[0]
    assert row[23:27] == ["recurring", None, "INR", "a" * 64]


def test_proforma_and_expense_report_are_not_invoices(client, users, fake_extract):
    for doc in ("proforma_invoice", "expense_report"):
        fake_extract.append(not_an_invoice().model_copy(update={"document_type": doc}))
        upload(client, users["priya"], doc.encode())
    run_worker()
    assert [get(client, i)["status"] for i in (1, 2)] == [Status.NOT_INVOICE, Status.NOT_INVOICE]


def test_invalid_correction_is_422_not_500(client, users, fake_extract):
    fake_extract.append(make_extracted(grand_total=1))
    inv = upload(client, users["priya"])["invoice_id"]
    run_worker()
    r = client.patch(f"/api/v1/invoices/{inv}", json={"invoice_kind": None}, headers=h(users["priya"]))
    assert r.status_code == 422 and "invoice_kind" in r.json()["detail"]
    r = client.patch(f"/api/v1/invoices/{inv}", json={"invoice_kind": "weekly"}, headers=h(users["priya"]))
    assert r.status_code == 422


# ---------- AI service down (internet or OpenAI outage)

def test_ai_outage_does_not_use_up_attempts(client, users, fake_extract):
    """Real case: Agno returns 'Connection error.' when OpenAI is unreachable; before, every queued
    invoice burnt its 3 attempts within seconds and ended FAILED."""
    fake_extract += [extract.AIServiceUnavailable("AI service unavailable: Connection error.")] * 5 + [make_extracted()]
    inv = upload(client, users["priya"])["invoice_id"]
    for _ in range(5):                                   # five outages in a row
        assert worker.process(worker.claim_next()) is False
        d = get(client, inv)
        assert d["status"] == Status.RECEIVED and d["error"].startswith("Waiting for the AI service")
    assert worker.process(worker.claim_next()) is True   # service back: processed normally
    run_worker()
    assert get(client, inv)["status"] == Status.EXPORTED


def test_agno_connection_error_text_becomes_ai_unavailable(tmp_path, monkeypatch):
    pdf = tmp_path / "a.pdf"
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "TAX INVOICE total 6136.00 " * 5)
    doc.save(pdf)
    doc.close()
    err = SimpleNamespace(content="Connection error.", metrics=None, status="RunStatus.error")
    monkeypatch.setattr(extract.extractor, "run", lambda *a, **k: err)
    with pytest.raises(extract.AIServiceUnavailable):
        extract.extract_invoice(str(pdf))
    bad_key = SimpleNamespace(content="Error code: 401 - Your API key has been invalidated.", metrics=None,
                              status="RunStatus.error")
    monkeypatch.setattr(extract.extractor, "run", lambda *a, **k: bad_key)
    with pytest.raises(extract.AIAccountProblem):        # the account, not the document: it waits
        extract.extract_invoice(str(pdf))
    # The model's own (unparseable) text mentioning "INV/500" or "timeout" is not an outage:
    # retrying it as one would block the queue forever
    garbled = SimpleNamespace(content="Invoice INV/500, payment timeout 30 days", metrics=None,
                              status="RunStatus.completed")
    monkeypatch.setattr(extract.extractor, "run", lambda *a, **k: garbled)
    with pytest.raises(ValueError):
        extract.extract_invoice(str(pdf))



def test_no_openai_credits_documents_wait_instead_of_failing(client, users, fake_extract, monkeypatch):
    """Seen for real: 'You have no credits remaining' failed every document after 3 attempts. The
    account isn't the document's fault: it waits in the queue (attempt not used) and the status
    endpoint explains why."""
    monkeypatch.setattr(worker, "AI_ACCOUNT_PAUSE_SECONDS", 0)
    inv_id = upload(client, users["priya"])["invoice_id"]
    for _ in range(5):   # five tries while the account is empty: never FAILED
        fake_extract.append(extract.AIAccountProblem("OpenAI account problem: You have no credits remaining."))
        assert worker.claim_next() == inv_id
        assert worker.process(inv_id) is False
    d = get(client, inv_id)
    assert d["status"] == Status.RECEIVED and "OpenAI account" in (d["error"] or "")
    st = client.get("/api/v1/system/status").json()
    assert st["ai"]["ok"] is False and st["ai"]["kind"] == "account" and st["ai"]["waiting_on_account"] == 1
    fake_extract.append(make_extracted())                   # credits added
    run_worker()
    assert get(client, inv_id)["status"] == Status.EXPORTED
    assert client.get("/api/v1/system/status").json()["ai"]["ok"] is True
