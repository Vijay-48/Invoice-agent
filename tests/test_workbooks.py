# ruff: noqa: F811  (pytest fixtures imported from test_flow)
"""One Excel workbook per financial year (D-96)."""
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook

from app import excel, worker
from app.config import settings
from app.models import Status
from tests.conftest import make_extracted
from tests.test_flow import client, fake_extract, get, run_worker, upload   # noqa: F401 (fixtures)


def _ids(fy):
    return [r[0] for r in load_workbook(excel.workbook_path(fy))["Invoices"].iter_rows(min_row=2, values_only=True)]


def test_financial_year_of_a_date():
    assert excel.fy_label(date(2026, 3, 31)) == "2025-26"
    assert excel.fy_label(date(2026, 4, 1)) == "2026-27"
    assert excel.fy_label(date(2099, 6, 1)) == "2099-00"
    assert excel.workbook_path("2026-27").name == "invoices FY2026-27.xlsx"


def test_invoices_go_to_the_workbook_of_their_invoice_date(client, users, fake_extract):
    fake_extract += [make_extracted(invoice_number="A-1", invoice_date="2026-03-31"),
                     make_extracted(invoice_number="A-2", invoice_date="2026-04-01")]
    march = upload(client, users["priya"], b"one")["invoice_id"]
    april = upload(client, users["priya"], b"two")["invoice_id"]
    run_worker()
    assert _ids("2025-26") == [march] and _ids("2026-27") == [april]
    assert [w["year"] for w in client.get("/api/v1/excel/workbooks").json()] == ["2026-27", "2025-26"]
    # Default download: this financial year's workbook; ?year= picks another
    r = client.get("/api/v1/excel")
    assert r.status_code == 200 and 'FY2026-27' in r.headers["content-disposition"]
    assert 'FY2025-26' in client.get("/api/v1/excel?year=2025-26").headers["content-disposition"]
    assert client.get("/api/v1/excel?year=2019-20").status_code == 404
    assert client.get("/api/v1/excel?year=26-27").status_code == 422


def test_a_workbook_open_in_excel_holds_back_only_its_year(client, users, fake_extract, monkeypatch):
    real_replace = excel.os.replace

    def replace(src, dst):
        if "FY2025-26" in str(dst):
            raise PermissionError("file is open")
        return real_replace(src, dst)

    monkeypatch.setattr(excel.os, "replace", replace)
    fake_extract += [make_extracted(invoice_number="B-1", invoice_date="2026-02-10"),
                     make_extracted(invoice_number="B-2", invoice_date="2026-09-15")]
    old = upload(client, users["priya"], b"one")["invoice_id"]
    new = upload(client, users["priya"], b"two")["invoice_id"]
    run_worker()
    assert get(client, new)["status"] == Status.EXPORTED
    waiting = get(client, old)
    assert waiting["status"] == Status.APPROVED and "FY2025-26" in waiting["error"]

    monkeypatch.setattr(excel.os, "replace", real_replace)
    worker.export()
    assert get(client, old)["status"] == Status.EXPORTED and _ids("2025-26") == [old]


def test_split_off_keeps_the_single_workbook(client, users, fake_extract, monkeypatch):
    monkeypatch.setattr(settings, "excel_per_financial_year", False)
    fake_extract.append(make_extracted())
    inv = upload(client, users["priya"])["invoice_id"]
    run_worker()
    assert Path(settings.excel_path).exists() and _ids(None) == [inv]
    assert client.get("/api/v1/excel").headers["content-disposition"].endswith('"invoices.xlsx"')


def test_workbook_from_before_the_split_is_kept_and_listed(client, users, fake_extract):
    Path(settings.excel_path).parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.active.title = "Invoices"
    wb.save(settings.excel_path)
    fake_extract.append(make_extracted())
    upload(client, users["priya"])
    run_worker()
    books = client.get("/api/v1/excel/workbooks").json()
    assert [(b["year"], b["file"]) for b in books] == [("2026-27", "invoices FY2026-27.xlsx"), (None, "invoices.xlsx")]
    assert client.get("/api/v1/excel?year=all").headers["content-disposition"].endswith('"invoices.xlsx"')
    assert load_workbook(settings.excel_path)["Invoices"].max_row == 1   # untouched
