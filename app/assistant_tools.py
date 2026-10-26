"""Tools the AI assistant can call. Each returns compact JSON text.

All tools only READ data (database, Excel workbook, email log) or create export files.
Changing an invoice is never done by a tool: `propose_action` only suggests it, and the
person confirms it with a button in the frontend, which calls the normal API.

Text that came from documents and emails (supplier names, subjects, bodies) is data written
by outsiders; the assistant's instructions say to never follow instructions found in it.
"""
import csv
import io
import json
import re
import statistics
import uuid
from collections import defaultdict
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from sqlalchemy import func, or_, select

from app.config import settings
from app.db import SessionLocal
from app.excel import _safe
from app.models import AuditLog, EmailMessage, Invoice, Status, User

MAX_ROWS = 50
API = "/api/v1"
STATUSES = [s for s in vars(Status) if s.isupper()]
ACTIONS = {   # what a person can do, from which statuses (mirrors app/api.py)
    "approve": {Status.NEEDS_REVIEW},
    "reject": {Status.NEEDS_REVIEW, Status.DUPLICATE, Status.NOT_INVOICE, Status.NEEDS_PASSWORD},
    "reprocess": {Status.NEEDS_REVIEW, Status.FAILED, Status.NOT_INVOICE, Status.DUPLICATE, Status.REJECTED,
                  Status.NEEDS_PASSWORD},
    "process-as-invoice": {Status.NOT_INVOICE},
}


def _json(data) -> str:
    return json.dumps(data, default=str, ensure_ascii=False)


def _num(v) -> float:
    try:
        return float(v or 0)
    except (TypeError, ValueError):
        return 0.0


def file_url(inv: Invoice) -> str:
    return f"{API}/invoices/{inv.id}/file"


def summary(inv: Invoice) -> dict:
    d = inv.extracted or {}
    return {"invoice_id": inv.id, "status": inv.status, "document": d.get("document_type"),
            "supplier": inv.supplier_name, "supplier_gstin": inv.supplier_gstin, "invoice_no": inv.invoice_no,
            "invoice_date": inv.invoice_date, "grand_total": _num(inv.grand_total),
            "currency": d.get("currency") or "INR", "source": inv.source, "file_name": inv.original_filename,
            "received": inv.created_at, "status_reason": (inv.status_reason or "")[:200] or None,
            "file_url": file_url(inv)}


def _filtered(session, text="", status="", supplier="", document_type="", date_from="", date_to="",
              min_total=None, max_total=None, source=""):
    q = select(Invoice)
    if text:
        like = f"%{text.strip()}%"
        q = q.where(or_(Invoice.supplier_name.ilike(like), Invoice.invoice_no.ilike(like),
                        Invoice.original_filename.ilike(like), Invoice.email_subject.ilike(like),
                        Invoice.sender.ilike(like), Invoice.supplier_gstin.ilike(like)))
    if status:
        wanted = [s.strip().upper() for s in status.split(",") if s.strip()]
        q = q.where(Invoice.status.in_(wanted))
    if supplier:
        q = q.where(or_(Invoice.supplier_name.ilike(f"%{supplier.strip()}%"),
                        Invoice.supplier_gstin.ilike(f"%{supplier.strip()}%")))
    if document_type:
        q = q.where(func.coalesce(Invoice.extracted["document_type"].astext, "") == document_type.strip())
    if date_from:
        q = q.where(Invoice.invoice_date >= date_from)
    if date_to:
        q = q.where(Invoice.invoice_date <= date_to)
    if min_total is not None:
        q = q.where(Invoice.grand_total >= min_total)
    if max_total is not None:
        q = q.where(Invoice.grand_total <= max_total)
    if source:
        q = q.where(Invoice.source == source.strip().lower())
    return q


# ---------- tools

def search_invoices(text: str = "", status: str = "", supplier: str = "", document_type: str = "",
                    date_from: str = "", date_to: str = "", min_total: float | None = None,
                    max_total: float | None = None, source: str = "", limit: int = 20) -> str:
    """Find invoices/documents. All filters are optional and combined.

    Args:
        text: words to look for in supplier name, invoice number, GSTIN, file name, email subject or sender.
        status: one or more of RECEIVED, PROCESSING, NEEDS_REVIEW, NEEDS_PASSWORD, APPROVED, EXPORTED,
            DUPLICATE, NOT_INVOICE, REJECTED, FAILED (comma-separated).
        supplier: part of the supplier name or GSTIN.
        document_type: tax_invoice, bill_of_supply, credit_note, debit_note, journal_entry, ...
        date_from: invoice date from, YYYY-MM-DD. date_to: invoice date until, YYYY-MM-DD.
        min_total: minimum grand total. max_total: maximum grand total.
        source: gmail or upload.
        limit: how many to return (max 50), newest first.
    Returns JSON with the number of matches and a summary of each (with a file_url to the original document).
    """
    with SessionLocal() as s:
        q = _filtered(s, text, status, supplier, document_type, date_from, date_to, min_total, max_total, source)
        total = s.scalar(select(func.count()).select_from(q.subquery()))
        rows = s.scalars(q.order_by(Invoice.created_at.desc(), Invoice.id.desc()).limit(min(max(limit, 1), MAX_ROWS)))
        return _json({"matches": total, "invoices": [summary(i) for i in rows]})


def get_invoice(invoice_id: int) -> str:
    """Everything about one invoice: all extracted values and line items, every check with its result,
    why it has its status, the email it came from, its processing history, and links to the original
    file. Use it before explaining an invoice or suggesting an action on it."""
    with SessionLocal() as s:
        inv = s.get(Invoice, invoice_id)
        if inv is None:
            return _json({"error": f"No invoice #{invoice_id}"})
        history = s.scalars(select(AuditLog).where(AuditLog.invoice_id == inv.id)
                            .order_by(AuditLog.created_at, AuditLog.id)).all()
        user = s.get(User, inv.user_id)
        email = None
        if inv.source == "gmail" and inv.source_ref:
            e = s.scalar(select(EmailMessage).where(EmailMessage.user_id == inv.user_id,
                                                    EmailMessage.message_id == inv.source_ref.split("/")[0]))
            if e:
                email = {"from": e.sender, "subject": e.subject, "received": e.received_at, "folders": e.folders,
                         "gmail_label": e.ap_label, "body_start": (e.body_snippet or "")[:300]}
        return _json({**summary(inv), "user": user.name if user else None, "attempts": inv.attempts,
                      "error": inv.error, "duplicate_of": inv.duplicate_of_id, "approved_by": inv.approved_by,
                      "approved_at": inv.approved_at, "exported_at": inv.exported_at,
                      "extracted": inv.extracted, "checks": inv.checks, "email": email,
                      "history": [{"at": h.created_at, "stage": h.stage, "by": h.actor, "message": h.message[:200]}
                                  for h in history][-25:]})


# Documents that don't add to what is owed: copies, rejected, not bills, unreadable
NOT_COUNTED = {Status.DUPLICATE, Status.REJECTED, Status.NOT_INVOICE, Status.FAILED}


def invoice_statistics(group_by: str = "status", date_from: str = "", date_to: str = "", status: str = "",
                       supplier: str = "") -> str:
    """Totals and counts grouped by one of: status, supplier, month, source, document_type, gst_rate.
    Amounts per group: count, taxable value, CGST, SGST, IGST, grand total (credit notes count negative).
    Duplicates (incl. journal vouchers of invoices already here), rejected, non-invoices and failed
    documents are left out unless `status` asks for them. Amounts in another currency than INR get
    their own groups (e.g. "Supplier X (USD)"). Optional filters: invoice date range (YYYY-MM-DD),
    status (comma-separated), supplier."""
    with SessionLocal() as s:
        rows = s.scalars(_filtered(s, status=status, supplier=supplier, date_from=date_from,
                                   date_to=date_to)).all()
    excluded = 0
    if not status:
        excluded = sum(1 for inv in rows if inv.status in NOT_COUNTED)
        rows = [inv for inv in rows if inv.status not in NOT_COUNTED]
    groups: dict = defaultdict(lambda: {"count": 0, "taxable": 0.0, "cgst": 0.0, "sgst": 0.0, "igst": 0.0,
                                        "grand_total": 0.0})
    for inv in rows:
        d = inv.extracted or {}
        sign = -1 if d.get("document_type") == "credit_note" else 1
        if group_by == "gst_rate":
            for line in d.get("lines") or []:
                g = groups[f"{_num(line.get('gst_rate')):g}%"]
                g["count"] += 1
                g["taxable"] += sign * _num(line.get("taxable_value"))
                g["grand_total"] += sign * _num(line.get("taxable_value")) * (1 + _num(line.get("gst_rate")) / 100)
            continue
        key = {"status": inv.status, "supplier": inv.supplier_name or "(unknown)", "source": inv.source,
               "month": (inv.invoice_date or "unknown")[:7], "document_type": d.get("document_type") or "unknown",
               }.get(group_by, inv.status)
        currency = (d.get("currency") or "INR").upper()
        g = groups[key if currency == "INR" else f"{key} ({currency})"]
        g["count"] += 1
        for k, field in (("taxable", "taxable_total"), ("cgst", "cgst"), ("sgst", "sgst"), ("igst", "igst"),
                         ("grand_total", "grand_total")):
            g[k] += sign * _num(d.get(field))
    out = [{"group": k, **{f: round(v, 2) if isinstance(v, float) else v for f, v in g.items()}}
           for k, g in sorted(groups.items(), key=lambda kv: -abs(kv[1]["grand_total"]))]
    return _json({"group_by": group_by, "documents": len(rows), "left_out_duplicates_rejected_etc": excluded,
                  "groups": out[:MAX_ROWS]})


def read_excel(sheet: str = "Invoices", search: str = "", limit: int = 30) -> str:
    """Read the Excel workbook of approved invoices (output/invoices.xlsx) exactly as it is on disk.

    Args:
        sheet: "Invoices" (one row per invoice) or "Line Items" (one row per item).
        search: only rows containing this text in any cell (case-insensitive).
        limit: rows to return (max 50).
    Returns headers, matching rows, totals of the amount columns over ALL matching rows, and a
    download_url for the workbook."""
    path = Path(settings.excel_path)
    if not path.exists():
        return _json({"error": "The Excel workbook doesn't exist yet: no invoice has been approved."})
    # Read from memory: an open file handle would stop the worker from replacing the workbook (Windows)
    wb = load_workbook(io.BytesIO(path.read_bytes()), read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        return _json({"error": f"No sheet '{sheet}'", "sheets": wb.sheetnames})
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    headers = [str(h) if h is not None else "" for h in (rows[0] if rows else [])]
    data = [r for r in rows[1:] if any(v is not None for v in r)]
    if search:
        needle = search.lower()
        data = [r for r in data if any(needle in str(v).lower() for v in r if v is not None)]
    money_cols = [i for i, h in enumerate(headers)
                  if h in ("Taxable Value", "CGST", "SGST", "IGST", "Round Off", "Grand Total", "GST Amount")]
    totals = {headers[i]: round(sum(_num(r[i]) for r in data if i < len(r)), 2) for i in money_cols}
    wb.close()

    def cell(v):   # Excel dates come back as datetimes at midnight
        if isinstance(v, datetime):
            return v.date().isoformat() if (v.hour, v.minute, v.second) == (0, 0, 0) else v.isoformat(" ", "minutes")
        return v
    return _json({"sheet": sheet, "headers": headers, "matching_rows": len(data), "totals": totals,
                  "rows": [dict(zip(headers, map(cell, r))) for r in data[:min(max(limit, 1), MAX_ROWS)]],
                  "download_url": f"{API}/excel"})


def list_emails(outcome: str = "", days: int = 7, limit: int = 30) -> str:
    """The Gmail email log: emails seen in the last `days` days and what happened to each attachment.
    outcome: done (documents queued), no_documents, skipped, failed, or empty for all."""
    since = datetime.now(timezone.utc) - timedelta(days=max(days, 1))
    with SessionLocal() as s:
        q = select(EmailMessage).where(EmailMessage.received_at >= since)
        if outcome:
            q = q.where(EmailMessage.outcome == outcome)
        rows = s.scalars(q.order_by(EmailMessage.received_at.desc()).limit(min(max(limit, 1), MAX_ROWS))).all()
        return _json({"emails": [{"received": e.received_at, "from": e.sender, "subject": e.subject,
                                  "folders": e.folders, "outcome": e.outcome, "error": e.error,
                                  "gmail_label": e.ap_label,
                                  "attachments": [{k: a.get(k) for k in ("name", "result", "invoice_id",
                                                                         "invoice_ids", "note")}
                                                  for a in e.attachments or []]} for e in rows]})


def needs_attention() -> str:
    """What needs a person now: invoices waiting for review (with reasons), protected PDFs waiting for a
    password, failed documents, approved invoices not yet in Excel, recent duplicates, emails that failed,
    and unusually large invoices (over 3x the supplier's usual amount)."""
    with SessionLocal() as s:
        def by_status(st, n=MAX_ROWS):
            return [summary(i) for i in s.scalars(select(Invoice).where(Invoice.status == st)
                                                  .order_by(Invoice.created_at.desc()).limit(n))]
        unusual = []
        totals = defaultdict(list)
        for gstin, total in s.execute(select(Invoice.supplier_gstin, Invoice.grand_total).where(
                Invoice.status == Status.EXPORTED, Invoice.supplier_gstin.is_not(None))):
            totals[gstin].append(_num(total))
        for inv in s.scalars(select(Invoice).where(Invoice.status.in_([Status.NEEDS_REVIEW, Status.APPROVED,
                                                                       Status.EXPORTED]),
                                                   Invoice.supplier_gstin.is_not(None))):
            past = [t for t in totals.get(inv.supplier_gstin, []) if t]
            if len(past) >= 3 and _num(inv.grand_total) > 3 * statistics.median(past):
                unusual.append({**summary(inv), "usual_amount": round(statistics.median(past), 2)})
        failed_emails = s.scalars(select(EmailMessage).where(EmailMessage.outcome == "failed")
                                  .order_by(EmailMessage.received_at.desc()).limit(10)).all()
        return _json({
            "needs_review": by_status(Status.NEEDS_REVIEW), "needs_password": by_status(Status.NEEDS_PASSWORD),
            "failed": by_status(Status.FAILED), "approved_not_in_excel_yet": by_status(Status.APPROVED),
            "recent_duplicates": by_status(Status.DUPLICATE, 10), "unusually_large": unusual[:20],
            "failed_emails": [{"received": e.received_at, "from": e.sender, "subject": e.subject, "error": e.error}
                              for e in failed_emails]})


def supplier_summary(supplier: str = "") -> str:
    """Per supplier: GSTIN, number of invoices, total amount in Excel (exported), amount waiting for review,
    first and last invoice date, and the statuses of their documents. `supplier`: part of a name or GSTIN
    to look at one supplier; empty for all (largest first)."""
    with SessionLocal() as s:
        rows = s.scalars(_filtered(s, supplier=supplier).where(Invoice.supplier_name.is_not(None))).all()
    out: dict = defaultdict(lambda: {"gstins": set(), "documents": 0, "exported_total": 0.0, "in_review_total": 0.0,
                                     "first_invoice": None, "last_invoice": None, "statuses": defaultdict(int)})
    for inv in rows:
        g = out[(inv.supplier_name or "").strip().upper()]
        g["name"] = inv.supplier_name
        if inv.supplier_gstin:
            g["gstins"].add(inv.supplier_gstin)
        g["documents"] += 1
        g["statuses"][inv.status] += 1
        if inv.status == Status.EXPORTED:
            g["exported_total"] += _num(inv.grand_total)
        if inv.status == Status.NEEDS_REVIEW:
            g["in_review_total"] += _num(inv.grand_total)
        if inv.invoice_date:
            g["first_invoice"] = min(filter(None, [g["first_invoice"], inv.invoice_date]))
            g["last_invoice"] = max(filter(None, [g["last_invoice"], inv.invoice_date]))
    result = [{**{k: v for k, v in g.items() if k not in ("gstins", "statuses")}, "gstins": sorted(g["gstins"]),
               "statuses": dict(g["statuses"]), "exported_total": round(g["exported_total"], 2),
               "in_review_total": round(g["in_review_total"], 2)} for g in out.values()]
    result.sort(key=lambda g: -(g["exported_total"] + g["in_review_total"]))
    return _json({"suppliers": result[:MAX_ROWS]})


def export_invoices(text: str = "", status: str = "", supplier: str = "", document_type: str = "",
                    date_from: str = "", date_to: str = "", min_total: float | None = None,
                    max_total: float | None = None, source: str = "", file_format: str = "xlsx") -> str:
    """Create a downloadable file (xlsx or csv) of the invoices matching the filters (same filters as
    search_invoices, no row limit). The xlsx has an Invoices sheet and a Line Items sheet.
    Returns the download_url to give to the person."""
    fmt = "csv" if file_format.lower().strip(". ") == "csv" else "xlsx"
    with SessionLocal() as s:
        rows = s.scalars(_filtered(s, text, status, supplier, document_type, date_from, date_to, min_total,
                                   max_total, source).order_by(Invoice.invoice_date, Invoice.id)).all()
    headers = ["Invoice ID", "Status", "Document", "Supplier", "Supplier GSTIN", "Buyer GSTIN", "Invoice No",
               "Invoice Date", "Taxable Value", "CGST", "SGST", "IGST", "Round Off", "Grand Total", "Currency",
               "Source", "File"]
    table, lines = [], []
    for inv in rows:
        d = inv.extracted or {}
        table.append([inv.id, inv.status, d.get("document_type"), inv.supplier_name, inv.supplier_gstin,
                      d.get("buyer_gstin"), inv.invoice_no, inv.invoice_date, _num(d.get("taxable_total")),
                      _num(d.get("cgst")), _num(d.get("sgst")), _num(d.get("igst")), _num(d.get("round_off")),
                      _num(d.get("grand_total")), d.get("currency") or "INR", inv.source, inv.original_filename])
        for n, l in enumerate(d.get("lines") or [], start=1):
            lines.append([inv.id, inv.invoice_no, n, l.get("description"), l.get("hsn_sac"), _num(l.get("quantity")),
                          _num(l.get("unit_price")), _num(l.get("taxable_value")), _num(l.get("gst_rate"))])
    folder = Path(settings.export_dir)
    folder.mkdir(parents=True, exist_ok=True)
    name = f"export-{date.today():%Y%m%d}-{uuid.uuid4().hex[:8]}.{fmt}"
    if fmt == "csv":
        with open(folder / name, "w", newline="", encoding="utf-8-sig") as fh:
            w = csv.writer(fh)
            w.writerow(headers)
            w.writerows([_safe(v) for v in r] for r in table)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Invoices"
        ws.append(headers)
        for r in table:
            ws.append([_safe(v) for v in r])
        ws2 = wb.create_sheet("Line Items")
        ws2.append(["Invoice ID", "Invoice No", "#", "Description", "HSN/SAC", "Qty", "Rate", "Taxable Value",
                    "GST %"])
        for r in lines:
            ws2.append([_safe(v) for v in r])
        for sheet in (ws, ws2):
            sheet.freeze_panes = "A2"
        wb.save(folder / name)
    return _json({"rows": len(table), "file_name": name, "download_url": f"{API}/assistant/files/{name}"})


def propose_action(invoice_id: int, action: str, reason: str) -> str:
    """Suggest an action on an invoice for the person to confirm with a button. NOTHING is changed by this
    tool: the person decides. action: approve, reject, reprocess (run the AI again) or process-as-invoice.
    Only propose an action that is allowed for the invoice's current status; explain why in `reason`."""
    action = action.strip().lower()
    with SessionLocal() as s:
        inv = s.get(Invoice, invoice_id)
        if inv is None:
            return _json({"error": f"No invoice #{invoice_id}"})
        if action not in ACTIONS:
            return _json({"error": f"Unknown action '{action}'", "allowed": sorted(ACTIONS)})
        if inv.status not in ACTIONS[action]:
            allowed = [a for a, st in ACTIONS.items() if inv.status in st]
            return _json({"error": f"Invoice #{invoice_id} is {inv.status}; '{action}' isn't possible",
                          "possible_actions": allowed})
        # `version`: the invoice as judged; the button refuses to act if it changed since (a saved chat's
        # button may be clicked days later)
        return _json({"proposal": {"invoice_id": invoice_id, "action": action, "reason": reason[:300],
                                   "invoice": summary(inv), "version": inv.updated_at.isoformat()}})


def gmail_overview() -> str:
    """Is the Gmail automation working: listener connection, last event and sync, connected inboxes."""
    from app import gmail_listener
    with SessionLocal() as s:
        inboxes = [{"user": u.name, "connected": bool(u.gmail_account_id), "last_sync": u.gmail_last_sync_at}
                   for u in s.scalars(select(User).order_by(User.id))]
        counts = dict(s.execute(select(EmailMessage.outcome, func.count()).group_by(EmailMessage.outcome)).all())
    return _json({"listener": gmail_listener.status, "inboxes": inboxes, "emails_by_outcome": counts})


TOOLS = [search_invoices, get_invoice, invoice_statistics, read_excel, list_emails, needs_attention,
         supplier_summary, export_invoices, propose_action, gmail_overview]

SAFE_FILE = re.compile(r"^export-\d{8}-[0-9a-f]{8}\.(xlsx|csv)$")
