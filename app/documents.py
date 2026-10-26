"""File types beyond PDF and images.

    Documents read directly   PDF, PNG/JPG/WEBP, Excel (.xlsx/.xlsm), CSV, Word (.docx)
    Containers unpacked       .zip and forwarded emails (.eml): each document inside is
                              processed on its own
    Email body                an invoice written in the email body becomes a PDF
"""
import csv
import email
import html
import io
import re
import threading
import zipfile
import zlib
from email import policy
from pathlib import Path

import pymupdf

IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp"}
TEXT_EXT = {".xlsx", ".xlsm", ".csv", ".docx"}          # read as text, no page images
DOCUMENT_EXT = {".pdf"} | IMAGE_EXT | TEXT_EXT
CONTAINER_EXT = {".zip", ".eml"}
MIME_EXT = {
    "application/pdf": ".pdf", "image/png": ".png", "image/jpeg": ".jpg", "image/webp": ".webp",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ".xlsx",
    "application/vnd.ms-excel.sheet.macroenabled.12": ".xlsm", "text/csv": ".csv",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
    "application/zip": ".zip", "application/x-zip-compressed": ".zip", "message/rfc822": ".eml",
}
MAX_TEXT_CHARS = 20000
MAX_CONTAINER_DEPTH = 2          # a zip inside a forwarded email is fine; deeper is ignored
MAX_FILES_PER_UPLOAD = 50        # documents one upload/attachment may unpack into (each costs an AI call)
UNPACK_BUDGET_FACTOR = 3         # total unpacked size: 3 x MAX_UPLOAD_MB per upload/attachment
LOGO_MAX_BYTES = 30_000          # images this small in an email are logos/signatures, not invoices
MAX_RENDER_PIXELS = 25_000_000   # one rendered page/crop (an A4 page at 200 dpi is ~4 million)


def render_dpi(rect, wanted: int) -> int:
    """The dpi to render `rect` (points) at without going over MAX_RENDER_PIXELS: a tiny PDF can
    declare a page kilometres wide, and rendering it at full resolution would take gigabytes."""
    pixels = (rect.width / 72 * wanted) * (rect.height / 72 * wanted)
    if pixels <= MAX_RENDER_PIXELS:
        return wanted
    return max(1, int(wanted * (MAX_RENDER_PIXELS / pixels) ** 0.5))


def kind(filename: str) -> str | None:
    ext = Path(filename).suffix.lower()
    if ext in DOCUMENT_EXT:
        return "document"
    if ext in CONTAINER_EXT:
        return "container"
    return None


def with_extension(filename: str, mime_type: str | None) -> str | None:
    """File name with a usable extension (from the MIME type if needed), or None if unsupported."""
    if kind(filename):
        return filename
    ext = MIME_EXT.get((mime_type or "").lower())
    return f"{filename}{ext}" if ext else None


class _Budget:
    """What one upload or attachment may unpack in total, across nested containers: a small zip
    can't expand into gigabytes (zip bomb) or thousands of AI calls."""
    def __init__(self, max_bytes: int):
        self.bytes_left = max_bytes * UNPACK_BUDGET_FACTOR
        self.files_left = MAX_FILES_PER_UPLOAD


def _is_logo(part) -> bool:
    """An image shown inside the email (signature, letterhead), not attached as a document."""
    if part.get_content_maintype() != "image":
        return False
    inline = part.get_content_disposition() == "inline" or (part.get_content_disposition() is None and part["Content-ID"])
    return bool(inline) or len(part.get_payload(decode=True) or b"") < LOGO_MAX_BYTES


def expand(filename: str, content: bytes, max_bytes: int, depth: int = 0,
           budget: _Budget | None = None) -> tuple[list[tuple[str, bytes]], list[str]]:
    """A document or container -> ([(name, bytes) of documents], [names skipped, with reason])."""
    budget = budget or _Budget(max_bytes)
    k = kind(filename)
    if k == "document":
        return [(filename, content)], []
    if k != "container" or depth >= MAX_CONTAINER_DEPTH:
        return [], [f"{filename} (not a supported file)"]
    inner: list[tuple[str, bytes]] = []
    skipped_parts: list[str] = []

    def take(name: str, size: int) -> bool:
        """Room left in the budget for one more file of this size?"""
        if budget.files_left <= 0:
            skipped_parts.append(f"{name} (more than {MAX_FILES_PER_UPLOAD} files)")
            return False
        if size > max_bytes or size > budget.bytes_left:
            skipped_parts.append(f"{name} (too large)")
            return False
        budget.files_left -= 1
        budget.bytes_left -= size
        return True

    if filename.lower().endswith(".zip"):
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as z:
                for info in z.infolist():
                    name = Path(info.filename).name
                    if info.is_dir() or name.startswith(".") or "__MACOSX" in info.filename:
                        continue
                    if not kind(name):
                        skipped_parts.append(f"{name} (not a supported file)")
                        continue
                    if not take(name, info.file_size):
                        continue
                    with z.open(info) as fh:   # never more than the declared (budgeted) size
                        inner.append((name, fh.read(info.file_size + 1)[:info.file_size]))
        except (zipfile.BadZipFile, RuntimeError, OSError, EOFError, zlib.error) as e:
            # RuntimeError: password-protected zip; zlib/EOF: damaged data
            return [], [f"{filename} (can't open: {e})"]
    else:   # .eml: a forwarded email
        msg = email.message_from_bytes(content, policy=policy.default)
        for part in msg.iter_attachments():
            name = part.get_filename() or "attachment"
            name = with_extension(name, part.get_content_type()) or name
            try:
                if _is_logo(part):
                    continue   # a signature/letterhead image: not a document, and not a reason to skip the body
                if part.get_content_type() == "message/rfc822":   # an email forwarded inside this one
                    data = part.get_content().as_bytes()
                else:
                    data = part.get_payload(decode=True) or b""
            except Exception as e:   # broken encoding: skip this part, keep the others
                skipped_parts.append(f"{name} (can't read: {type(e).__name__})")
                continue
            if data and take(name, len(data)):
                inner.append((name, data))
        body = msg.get_body(preferencelist=("plain", "html"))
        try:
            text = body.get_content() if body else ""
        except (LookupError, ValueError):   # unknown charset
            text = body.get_payload(decode=True).decode("utf-8", "replace") if body else ""
        if not inner and looks_like_invoice_text(text):
            inner.append((f"{Path(filename).stem} - body.pdf",
                          body_to_pdf(str(msg.get("subject", "")), str(msg.get("from", "")), text)))
    documents, skipped = [], [f"{filename}/{s}" for s in skipped_parts]
    for name, data in inner:
        if not data:
            skipped.append(name)
            continue
        docs, skip = expand(name, data, max_bytes, depth + 1, budget)
        documents += [(f"{filename}/{n}", d) for n, d in docs]
        skipped += [f"{filename}/{s}" for s in skip]
    return documents, skipped


# ---------- text documents (Excel, CSV, Word)

def document_text(path: str) -> str:
    """Readable text of an Excel/CSV/Word file for the AI (tables as tab-separated rows)."""
    ext = Path(path).suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        out = []
        for ws in wb.worksheets:
            out.append(f"--- sheet {ws.title} ---")
            for row in ws.iter_rows(values_only=True):
                cells = ["" if v is None else str(v) for v in row]
                if any(cells):
                    out.append("\t".join(cells).rstrip())
        text = "\n".join(out)
    elif ext == ".csv":
        raw = Path(path).read_bytes()
        decoded = raw.decode("utf-8-sig", errors="replace")
        text = "\n".join("\t".join(r) for r in csv.reader(io.StringIO(decoded)))
    elif ext == ".docx":
        import docx
        d = docx.Document(path)
        out = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            out += ["\t".join(c.text.strip() for c in row.cells) for row in t.rows]
        text = "\n".join(out)
    else:
        raise ValueError(f"Not a text document: {path}")
    return text[:MAX_TEXT_CHARS]


# ---------- invoices written in the email body

AMOUNT_RE = re.compile(r"(₹|rs\.?|inr)\s*[\d,]+(\.\d+)?|total\s*[:\-]?\s*[\d,]+\.\d{2}", re.IGNORECASE)
GSTIN_RE = re.compile(r"\b\d{2}[A-Z]{5}\d{4}[A-Z][1-9A-Z]Z[0-9A-Z]\b")


def looks_like_invoice_text(text: str) -> bool:
    """Body text that is itself an invoice: mentions invoice/bill and has a GSTIN or amounts."""
    if not text or len(text) < 80:
        return False
    has_word = re.search(r"\b(tax\s+)?invoice\b|\bbill\b", text, re.IGNORECASE)
    return bool(has_word and (GSTIN_RE.search(text) or len(AMOUNT_RE.findall(text)) >= 2))


# PyMuPDF (MuPDF) is not thread-safe: two threads using it at the same moment can crash the whole
# server. With several workers and Gmail threads, every PyMuPDF call goes through this lock. The calls
# are short; the slow work (the AI call, QR decoding) happens outside it.
PDF_LOCK = threading.RLock()


def body_to_pdf(subject: str, sender: str, text: str) -> bytes:
    """Render an email body as a PDF so it goes through the normal document pipeline."""
    with PDF_LOCK:
        return _body_to_pdf(subject, sender, text)


def _body_to_pdf(subject: str, sender: str, text: str) -> bytes:
    doc = pymupdf.open()
    content = (f"<p><b>Email from:</b> {html.escape(sender)}<br><b>Subject:</b> {html.escape(subject)}</p>"
               f"<pre style='font-family:sans-serif;white-space:pre-wrap'>{html.escape(text[:MAX_TEXT_CHARS])}</pre>")
    rect = pymupdf.Rect(36, 36, 559, 806)
    story = pymupdf.Story(html=content)
    writer = pymupdf.DocumentWriter(buf := io.BytesIO())
    more = True
    while more:
        dev = writer.begin_page(pymupdf.paper_rect("a4"))
        more, _ = story.place(rect)
        story.draw(dev)
        writer.end_page()
    writer.close()
    doc.close()
    return buf.getvalue()
